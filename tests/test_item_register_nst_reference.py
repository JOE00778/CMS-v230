import csv
import io
from unittest.mock import MagicMock
import openpyxl
import pytest
from data_warehouse.templates import item_entry_form as F, nst_item_master as N
from shared import item_register as R
from shared.nst_choices import Choices
from tests.test_item_register import _enrich, _row
CH=Choices(suppliers=('0001 TEST',),middle_by_large={'生活雑貨':('日用品雑貨',),'食品':('嗜好品（茶・コーヒー）',)})


def test_export_only_fallback_keeps_real_barcode_empty():
    rows=[{'型番':'ETERNAL-BONSAI-NV','JANコード':'','アイテム名':'盆栽'}]
    data=list(csv.DictReader(io.StringIO(R.build_nst_registration_csv(rows).decode('utf-8-sig'))))
    assert data[0]['JANコード']=='ETERNAL-BONSAI-NV'
    assert all(c in data[0] for c in ('税率','カートン高さ(cm)','カートン幅(cm)','カートン奥行(cm)','SKUID'))
    assert rows[0]['JANコード']==''
    conn=MagicMock();R.write_register_log(conn,rows,None)
    assert conn.executemany.call_args.args[1][0][1]==''


def test_carton_editing_and_prompts_for_all_products():
    nst,extra,_,_=_enrich([_row('4901234567894')]);frame=R.to_frame(nst,extra)
    fields=('カートン高さ(cm)','カートン幅(cm)','カートン奥行(cm)')
    assert all(c in frame[0] and c in R.EDITABLE for c in (*fields,'SKUID','税率','大分類','中分類'))
    issues=R.registration_issues(nst)
    assert all(any(c in i.message and '手入力' in i.message for i in issues) for c in fields)
    for c,v in zip(fields,('30','40','50')):frame[0][c]=v
    out,_,_,_=R.from_frame(frame,nst,extra)
    assert [out[0][c] for c in fields]==['30','40','50']
    assert not any('カートン' in i.message for i in R.registration_issues(out))


def test_tea_review_and_tax_recalculation_after_edit():
    nst=[{'型番':'TEA','JANコード':'4973028760248','アイテム名':'抹茶入り玄米茶6個入り','大分類':'生活雑貨','中分類':'日用品雑貨','仕入区分':'輸出専用','商品原価':'100','仕入先1：仕入先':'0001 TEST','商品担当者':'043 徐越'}]
    frame=R.to_frame(nst,{})
    assert any('分類・税率' in i.message for i in R.registration_issues(nst))
    frame[0].update({'大分類':'食品','中分類':'嗜好品（茶・コーヒー）','税率':'★'})
    out,_,_,_=R.from_frame(frame,nst,{},choices=CH)
    assert out[0]['納税スケジュール']=='仕入8%＆販売免税'
    assert not any('分類・税率' in i.message for i in R.registration_issues(out,choices=CH))
    frame[0]['中分類']='日用品雑貨'
    assert any(i.level=='error' and '中分類' in i.message for i in R.from_frame(frame,nst,{},choices=CH)[3])


def test_schema_71_and_applicable_conditions():
    assert [r['項目'] for r in R.nst_field_rules()]==[c for _,c in N.NST_MASTER_COLUMNS]
    export=R.registration_issues([{'型番':'X','仕入区分':'輸出専用'}])
    assert any('SKUID' in i.message and '確認' in i.message for i in export)
    assert not any('20GP入数' in i.message for i in export)
    assert any('20GP入数' in i.message for i in R.registration_issues([{'型番':'X','仕入区分':'海外仕入'}]))
    assert any('参考販売価格' in i.message for i in R.registration_issues([{'型番':'X','ライセンス率':'6'}]))


def test_template_full_fields_and_decimal_carton_text_lot():
    for c in ('カートン高さ(cm)','カートン幅(cm)','カートン奥行(cm)','SKUID','税率','メーカー型番','商品ランク'):assert c in F.TEMPLATE_COLUMNS
    wb=openpyxl.load_workbook(io.BytesIO(F.build_template(CH)));ws=wb[F.SHEET_INPUT]
    r={'型番':'SKU','アイテム名（任意）':'盆栽','商品原価':'100','仕入先1：仕入先':'0001 TEST','商品担当者':'043 徐越','大分類':'生活雑貨','中分類':'日用品雑貨','カートン入数':'1.5','発注ロット':'1色1000','カートン高さ(cm)':'３０','カートン幅(cm)':'40','カートン奥行(cm)':'50','SKUID':'FLAG-1','メーカー型番':'M-1'}
    for j,c in enumerate(F.TEMPLATE_COLUMNS,1):ws.cell(2,j,r.get(c,''))
    b=io.BytesIO();wb.save(b);rows,issues=F.read_upload(b.getvalue(),CH)
    assert len(rows)==1 and not any(i.level=='error' for i in issues)
    assert rows[0]['カートン高さ(cm)']=='30' and rows[0]['SKUID']=='FLAG-1' and rows[0]['発注ロット']=='1色1000'


def test_reference_sheet_skips_instructions_and_examples():
    wb=openpyxl.Workbook();ws=wb.active;ws.title='NetSuiteマスタ登録'
    ws.append([m for m,_ in N.NST_MASTER_COLUMNS]);ws.append([c for _,c in N.NST_MASTER_COLUMNS]);ws.append(['instructions']);ws.append(['example'])
    r={'型番':'SKU','JANコード':'SKU','アイテム名':'盆栽','商品原価':'100','仕入先1：仕入先':'0001 TEST','商品担当者':'043 徐越','大分類':'生活雑貨','中分類':'日用品雑貨','カートン入数':'24','発注ロット':'12'}
    ws.append([r.get(c,'') for _,c in N.NST_MASTER_COLUMNS]);b=io.BytesIO();wb.save(b);rows,issues=F.read_upload(b.getvalue(),CH)
    assert len(rows)==1 and not any(i.level=='error' for i in issues)
    assert rows[0]['JANコード']=='' and rows[0]['型番']=='SKU' and rows[0][F.COL_NAME_MANUAL]=='盆栽'


@pytest.mark.parametrize('value',['0','-1','abc','nan','inf'])
def test_invalid_carton_is_not_defaulted(value):
    nst,extra,_,_=_enrich([_row('4901234567894')]);frame=R.to_frame(nst,extra);frame[0]['カートン高さ(cm)']=value
    out,_,_,issues=R.from_frame(frame,nst,extra)
    assert not out[0].get('カートン高さ(cm)')
    assert any('カートン高さ(cm)' in i.message for i in issues)


def test_actual_16_products_end_to_end_and_review_after_simulated_corrections():
    import json
    from pathlib import Path
    from shared.hs_classify import Customs
    from shared.jan_web import JanMeta
    from shared.nst_choices import FALLBACK_MIDDLE_BY_LARGE
    original=json.loads((Path(__file__).parent/'fixtures/nst_registration_16.json').read_text())
    choices=Choices(suppliers=tuple(dict.fromkeys(r['仕入先1：仕入先'] for r in original)),middle_by_large=dict(FALLBACK_MIDDLE_BY_LARGE))
    wb=openpyxl.Workbook();ws=wb.active;ws.title=F.SHEET_INPUT;ws.append(F.TEMPLATE_COLUMNS)
    for r in original:
        row={**r,F.COL_NAME_MANUAL:r['アイテム名'],F.COL_MAKER_MANUAL:r.get('メーカー名','')}
        ws.append([row.get(c,'') for c in F.TEMPLATE_COLUMNS])
    b=io.BytesIO();wb.save(b);inputs,upload_issues=F.read_upload(b.getvalue(),choices)
    assert len(inputs)==16 and not any(i.level=='error' for i in upload_issues)
    metadata=MagicMock(side_effect=lambda jan:JanMeta(jan,'unused','','','ok'))
    weight=MagicMock(side_effect=AssertionError('no web queries'))
    nst,extra,_,_=R.enrich(inputs,choices=choices,prefix_makers={},fetch_meta=metadata,fetch_weight=weight,classify_customs=lambda *a,**k:Customs(None,None,None,'none'),want_weight=False,interval=0)
    for r in original:
        extra[r['型番']].update({'_hs':r['_hs'],'_name_en':r['_name_en']})
    assert metadata.call_count==10 and not weight.called
    frame=R.to_frame(nst,extra);out,bm,_,errors=R.from_frame(frame,nst,extra,choices=choices)
    assert len(out)==len(bm)==16 and not any(i.level=='error' for i in errors)
    checks=R.registration_issues(out,choices=choices)
    assert len({i.jan for i in checks if 'カートン高さ' in i.message})==16
    assert len({i.jan for i in checks if 'SKUID' in i.message})==16
    assert len({i.jan for i in checks if '分類・税率' in i.message})==5
    assert sum(not r.get('メーカー名') for r in out)==11
    csv_rows=list(csv.DictReader(io.StringIO(R.build_nst_registration_csv(out).decode('utf-8-sig'))))
    assert len(csv_rows)==16
    for a,b in zip(original,csv_rows):
        assert b['JANコード']==(a['JANコード'] or a['型番'])
        assert all(b[c]==a[c] for c in R._WEIGHT_FIELDS)
        assert all(not b[c] for c in F.COL_CARTON_DIMS)
    assert sum(not r.get('JANコード') for r in out)==6
    tea={i.jan for i in checks if '分類・税率' in i.message}
    # Synthetic correction round proves editing/recalculation, never supplies real carton facts.
    for rec in frame:
        rec.update({c:'30' for c in F.COL_CARTON_DIMS});rec['SKUID']='TEST-FLAG';rec['商品ランク']=''
        if rec['型番'] in tea:rec.update({'大分類':'食品','中分類':'嗜好品（茶・コーヒー）','税率':'★'})
    corrected,_,_,issues=R.from_frame(frame,nst,extra,choices=choices)
    assert len(corrected)==16 and not any(i.level=='error' for i in issues)
    assert R.registration_issues(corrected,choices=choices)==[]
    assert all(r['納税スケジュール']=='仕入8%＆販売免税' for r in corrected if r['型番'] in tea)
    assert all(not r.get(c) for r in out for c in F.COL_CARTON_DIMS)


def test_license_zero_does_not_require_license_price():
    checks=R.registration_issues([{'型番':'X','ライセンス率':'0','ライセンス料(AT)':'0'}])
    assert not any('参考販売価格' in i.message for i in checks)


def test_native_date_cell_is_read_in_reference_format():
    import datetime
    wb=openpyxl.load_workbook(io.BytesIO(F.build_template(CH)));ws=wb[F.SHEET_INPUT]
    r={'型番':'SKU','アイテム名（任意）':'盆栽','商品原価':'100','仕入先1：仕入先':'0001 TEST','商品担当者':'043 徐越','大分類':'生活雑貨','中分類':'日用品雑貨','カートン入数':'24','発売日':datetime.datetime(2026,10,14)}
    for j,c in enumerate(F.TEMPLATE_COLUMNS,1):ws.cell(2,j,r.get(c,''))
    b=io.BytesIO();wb.save(b);rows,issues=F.read_upload(b.getvalue(),CH)
    assert len(rows)==1 and not any(i.level=='error' for i in issues)
    assert rows[0]['発売日']=='2026/10/14'


def test_16_product_page_zip_and_checklist_use_only_mock_services():
    import json
    import zipfile
    from pathlib import Path
    from unittest.mock import patch
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    from shared import nst_choices as NC, jan_web, hs_ai
    original=json.loads((Path(__file__).parent/'fixtures/nst_registration_16.json').read_text())
    choices=NC.Choices(suppliers=tuple(dict.fromkeys(r['仕入先1：仕入先'] for r in original)),middle_by_large=dict(NC.FALLBACK_MIDDLE_BY_LARGE))
    wb=openpyxl.Workbook();ws=wb.active;ws.title=F.SHEET_INPUT;ws.append(F.TEMPLATE_COLUMNS)
    for r in original:
        row={**r,F.COL_NAME_MANUAL:r['アイテム名'],F.COL_MAKER_MANUAL:r.get('メーカー名','')};ws.append([row.get(c,'') for c in F.TEMPLATE_COLUMNS])
    b=io.BytesIO();wb.save(b);up=MagicMock();up.getvalue.return_value=b.getvalue();conn=MagicMock();conn.__enter__.return_value=conn
    st.cache_data.clear()
    with patch('shared.db.get_connection',return_value=conn),patch('streamlit.file_uploader',return_value=up),patch.object(NC,'load_choices',return_value=choices),patch.object(NC,'load_prefix_makers',return_value={}),patch.object(NC,'existing_jans',return_value=(set(),[])),patch.object(NC,'existing_item_codes',return_value=(set(),[])),patch.object(hs_ai,'is_configured',return_value=False),patch.object(jan_web,'fetch_meta',side_effect=lambda j:jan_web.JanMeta(j,'unused','','','ok')),patch.object(jan_web,'fetch_weight') as fw:
        at=AppTest.from_file(str(Path(__file__).parents[1]/'pages/15_📝_商品登录.py'),default_timeout=30);at.session_state['__auth_ok']=True
        at.run();at.button(key='page15_run').click().run()
        assert not at.exception,[e.value for e in at.exception]
        assert [m.value for m in at.metric][:2]==['16','0']
        editor=next(d.value for d in at.dataframe if 'カートン高さ(cm)' in d.value.columns)
        assert len(editor)==16 and editor['カートン高さ(cm)'].eq('').all()
        check=next(d.value for d in at.dataframe if '型番 / JAN' in d.value.columns and any('SKUID' in str(v) for v in d.value.values.ravel()))
        assert sum('カートン高さ' in str(v) for v in check.values.ravel())==16
        at.button(key='page15_gen').click().run()
        assert not at.exception,[e.value for e in at.exception]
        assert at.session_state['page15_zip_n']==16 and at.session_state['page15_pending_n']==16 and not fw.called
        z=zipfile.ZipFile(io.BytesIO(at.session_state['page15_zip_bytes']))
        assert '登録前の確認事項.csv' in z.namelist()
        nst_name=next(n for n in z.namelist() if 'NetSuite' in n)
        csv_rows=list(csv.DictReader(io.StringIO(z.read(nst_name).decode('utf-8-sig'))))
        assert len(csv_rows)==16 and all(r['JANコード']==(o['JANコード'] or o['型番']) for r,o in zip(csv_rows,original))
        assert all(not r[c] for r in csv_rows for c in F.COL_CARTON_DIMS)
        # Real barcode in the simulated ledger remains blank for six no-JAN models.
        assert sum(not p[1] for p in conn.executemany.call_args.args[1])==6


def test_reference_sheet_does_not_silently_change_procurement_mode():
    wb=openpyxl.Workbook();ws=wb.active;ws.title='NetSuiteマスタ登録'
    ws.append([m for m,_ in N.NST_MASTER_COLUMNS]);ws.append([c for _,c in N.NST_MASTER_COLUMNS]);ws.append(['instructions']);ws.append(['example'])
    r={'型番':'SKU','アイテム名':'盆栽','商品原価':'100','仕入先1：仕入先':'0001 TEST','商品担当者':'043 徐越','大分類':'生活雑貨','中分類':'日用品雑貨','カートン入数':'24','仕入区分':'海外仕入'}
    ws.append([r.get(c,'') for _,c in N.NST_MASTER_COLUMNS]);b=io.BytesIO();wb.save(b);rows,issues=F.read_upload(b.getvalue(),CH)
    assert rows==[]
    assert any(i.level=='error' and '仕入区分' in i.message and '輸出専用' in i.message for i in issues)


@pytest.mark.parametrize(('field','bad','good'), [('生原価','53','RMB53'),('中国手数料','5.3','5.3%'),('アイテム名','品Ⅰ','品1')])
def test_reference_format_rules_are_enforced(field,bad,good):
    assert any(field in message for message in F.validate_fields({field:bad}))
    assert F.validate_fields({field:good})==[]


def test_overseas_zero_amount_is_present_not_missing():
    row={'型番':'X','仕入区分':'海外仕入','開発費単価':0,'不良率自社負担単価':0,'関税(AT)':0}
    assert not any(any(f in i.message for f in ('開発費単価','不良率自社負担単価','関税(AT)')) for i in R.registration_issues([row]))


def test_registration_csv_exactly_matches_all_71_reference_columns():
    row={'型番':'ETERNAL-BONSAI-NV','JANコード':'','アイテム名':'盆栽'}
    parsed=list(csv.reader(io.StringIO(R.build_nst_registration_csv([row]).decode('utf-8-sig'))))
    assert parsed[0]==[c for _,c in N.NST_MASTER_COLUMNS]
    assert len(parsed[0])==len(parsed[1])==71
    assert parsed[1][2]=='ETERNAL-BONSAI-NV'
    assert parsed[1][27:30]==['','','']
    assert row['JANコード']==''
