"""page15（商品登录）を Streamlit AppTest で無頭実行 — 例外なく 3 タブとも描画されること。

load_choices / SuiteQL / DB はモック（网络/DB 不要）。
"""
from __future__ import annotations

import pathlib
from unittest import mock

import streamlit as st
from streamlit.testing.v1 import AppTest

from shared import nst_choices as NC

PAGE = pathlib.Path(__file__).resolve().parents[1] / "pages" / "15_📝_商品登录.py"


def test_page15_renders_all_tabs_without_exception():
    st.cache_data.clear()
    choices = NC.Choices(suppliers=("0001 ベストアンサ株式会社",),
                         middle_by_large=dict(NC.FALLBACK_MIDDLE_BY_LARGE),
                         warnings=("NST 分類の取得に失敗したため 2026-10-02 時点の定数を使用: mock",))
    with mock.patch("shared.db.get_connection", mock.MagicMock()), \
            mock.patch.object(NC, "load_choices", return_value=choices) as lc, \
            mock.patch.object(NC, "load_prefix_makers", return_value={}), \
            mock.patch.object(NC, "existing_jans", return_value=(set(), [])):
        at = AppTest.from_file(str(PAGE), default_timeout=60)
        at.session_state["__auth_ok"] = True
        at.run()
    assert not at.exception, [e.value for e in at.exception]
    assert len(at.tabs) == 3
    assert lc.called
    assert any("定数を使用" in w.value for w in at.warning)          # choices.warnings を表示
    labels = [b.label for b in at.button]
    assert any("処理する" in x for x in labels)                     # 🆕 商品登録
    assert any("查询 PG" in x for x in labels)                      # 📦 セット品登録
    assert not any("JD" in x for x in labels)                       # JD UI は削除済み
    assert at.checkbox[0].value is False                            # 網調べは既定 OFF


def _filled_upload(choices) -> bytes:
    import io

    from openpyxl import load_workbook

    from data_warehouse.templates import item_entry_form as FORM
    wb = load_workbook(io.BytesIO(FORM.build_template(choices)))
    ws = wb[FORM.SHEET_INPUT]
    vals = {"JANコード": "4901234567894", "商品原価": "100",
            "仕入先1：仕入先": "0001 ベストアンサ株式会社", "商品担当者": "043 徐越",
            "大分類": "食品", "中分類": "菓子", "カートン入数": "24", "発注ロット": "1",
            "パッケージ高さ(cm)": "10", "パッケージ幅(cm)": "5", "パッケージ奥行(cm)": "3"}
    for j, cell in enumerate(ws[1], start=1):
        if cell.value in vals:
            ws.cell(row=2, column=j, value=vals[cell.value])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_page15_process_and_generate_flow():
    """UL → ▶ 処理する → 📦 生成 まで。jancode はモック、台帳 INSERT は MagicMock conn。"""
    from shared import jan_web

    st.cache_data.clear()
    choices = NC.Choices(suppliers=("0001 ベストアンサ株式会社",),
                         middle_by_large=dict(NC.FALLBACK_MIDDLE_BY_LARGE))
    up = mock.MagicMock()
    up.getvalue.return_value = _filled_upload(choices)
    meta = jan_web.JanMeta("4901234567894", "テスト クッキー 100g", "テスト製菓株式会社", "", "ok")
    conn = mock.MagicMock()
    conn.__enter__.return_value = conn
    with mock.patch("shared.db.get_connection", return_value=conn), \
            mock.patch("streamlit.file_uploader", return_value=up), \
            mock.patch.object(NC, "load_choices", return_value=choices), \
            mock.patch.object(NC, "load_prefix_makers", return_value={}), \
            mock.patch.object(NC, "existing_jans", return_value=(set(), [])), \
            mock.patch.object(jan_web, "fetch_meta", return_value=meta), \
            mock.patch.object(jan_web, "INTERVAL", 0):
        at = AppTest.from_file(str(PAGE), default_timeout=60)
        at.session_state["__auth_ok"] = True
        at.run()
        at.button(key="page15_run").click().run()
        assert not at.exception, [e.value for e in at.exception]
        assert [m.value for m in at.metric][:2] == ["1", "0"]       # OK=1 NG=0
        at.button(key="page15_gen").click().run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.session_state["page15_zip_n"] == 1
    assert at.session_state["page15_log_msg"][0] == "ok"
    assert conn.executemany.called                                   # 台帳へ INSERT

    import csv
    import io
    import zipfile

    from openpyxl import load_workbook
    zf = zipfile.ZipFile(io.BytesIO(at.session_state["page15_zip_bytes"]))
    csv_name = next(n for n in zf.namelist() if n.endswith(".csv"))
    header, row = list(csv.reader(io.StringIO(zf.read(csv_name).decode("utf-8-sig"))))[:2]
    rec = dict(zip(header, row))
    assert header[0] == "型番" and rec["型番"] == "4901234567894"
    assert rec["アイテム名"] == "テスト クッキー 100g" and rec["税率"] == "★"
    assert rec["仕入先1：仕入先"] == "0001 ベストアンサ株式会社" and rec["仕入区分"] == "輸出専用"
    ws = load_workbook(io.BytesIO(zf.read(next(n for n in zf.namelist()
                                              if n.startswith("BM_"))))).active
    bm = dict(zip([c.value for c in ws[2]], [c.value for c in ws[3]]))
    assert bm["中文名称"] == "テスト クッキー 100g"
    assert len(bm["英文名称"] or "") <= 76


def _run_page(choices, up, *, meta, exist=frozenset(), conn=None):
    """▶ 処理する まで進めた AppTest と patch の ExitStack を返す（呼び出し側で close）。"""
    import contextlib

    from shared import jan_web
    conn = conn or mock.MagicMock()
    conn.__enter__.return_value = conn
    fm = meta if callable(meta) else mock.MagicMock(return_value=meta)
    stack = contextlib.ExitStack()
    for p in (mock.patch("shared.db.get_connection", return_value=conn),
              mock.patch("streamlit.file_uploader", return_value=up),
              mock.patch.object(NC, "load_choices", return_value=choices),
              mock.patch.object(NC, "load_prefix_makers", return_value={}),
              mock.patch.object(NC, "existing_jans", return_value=(set(exist), [])),
              mock.patch.object(jan_web, "fetch_meta", fm),
              mock.patch.object(jan_web, "INTERVAL", 0)):
        stack.enter_context(p)
    st.cache_data.clear()
    at = AppTest.from_file(str(PAGE), default_timeout=60)
    at.session_state["__auth_ok"] = True
    at.run()
    at.button(key="page15_run").click().run()
    return at, stack, fm, conn


def _choices():
    return NC.Choices(suppliers=("0001 ベストアンサ株式会社",),
                      middle_by_large=dict(NC.FALLBACK_MIDDLE_BY_LARGE))


def test_page15_existing_jan_is_excluded():
    from shared import jan_web
    ch = _choices()
    up = mock.MagicMock()
    up.getvalue.return_value = _filled_upload(ch)
    meta = jan_web.JanMeta("4901234567894", "テスト クッキー", "", "", "ok")
    at, stack, fm, _ = _run_page(ch, up, meta=meta, exist={"4901234567894"})
    with stack:
        assert not at.exception, [e.value for e in at.exception]
        assert [m.value for m in at.metric][:2] == ["0", "1"]       # OK=0 NG=1
        assert not fm.called
        assert any("登録済み" in str(v) for df in at.dataframe for v in df.value.values.ravel())


def test_page15_rerun_retries_after_jancode_error():
    from shared import jan_web
    ch = _choices()
    up = mock.MagicMock()
    up.getvalue.return_value = _filled_upload(ch)
    err = jan_web.JanMeta("4901234567894", "", "", "", "error", "http=403")
    at, stack, fm, _ = _run_page(ch, up, meta=err)
    with stack:
        assert fm.call_count == 1 and [m.value for m in at.metric][:2] == ["0", "1"]
        fm.return_value = jan_web.JanMeta("4901234567894", "テスト クッキー", "", "", "ok")
        at.button(key="page15_run").click().run()
        assert fm.call_count == 2 and [m.value for m in at.metric][:2] == ["1", "0"]


def test_page15_broken_xlsx_shows_error_and_other_tabs_render():
    ch = _choices()
    up = mock.MagicMock()
    up.getvalue.return_value = b"not an xlsx"
    at, stack, _, _ = _run_page(ch, up, meta=None)
    with stack:
        assert not at.exception, [e.value for e in at.exception]
        assert any("処理に失敗" in e.value for e in at.error)
        assert any(b.key == "page15_bundle_query" for b in at.button)


def test_page15_edit_hides_stale_zip_and_log_not_duplicated():
    from shared import jan_web
    ch = _choices()
    up = mock.MagicMock()
    up.getvalue.return_value = _filled_upload(ch)
    meta = jan_web.JanMeta("4901234567894", "テスト クッキー", "", "", "ok")
    at, stack, _, conn = _run_page(ch, up, meta=meta)
    with stack:
        at.button(key="page15_gen").click().run()
        at.button(key="page15_gen").click().run()
        assert conn.executemany.call_count == 1                     # 同じ内容は台帳に 1 回だけ
        assert any("page15_zip_dl" in b.proto.id for b in at.get("download_button"))
        # 生成後に表を編集 → 古い ZIP は伏せる
        at.session_state["page15_zip_src"] = "stale"
        at.run()
        assert not any("page15_zip_dl" in b.proto.id for b in at.get("download_button"))
        assert any("表を編集しました" in w.value for w in at.warning)
