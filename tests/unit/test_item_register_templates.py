"""商品登録 ZIP（NST csv / 斑马 xlsx）のテンプレ生成の回帰テスト。

現場から同じ不具合が二人ずつ挙がっていた（2026-08-18 修正）:
  #27 隋艶偉さん / #30 川崎さん … NST CSV から「サポート提供」列が消えて取込エラー
  #28 隋艶偉さん / #31 川崎さん … JD シート 1 列目（貨主ID）に値が入って取込エラー
  （JD 出力は v2 で廃止 · 2026-10-03。#28/#31 のテストも削除）

どちらも「テンプレ原文の列名に前後空白がある」「既定値へフォールバックする」という
静かな挙動が原因で、動かしてみるまで気付けなかった。ここで固定する。
"""
import csv
import io

from data_warehouse.templates import jd_bm_item_master as JBM
from data_warehouse.templates import nst_item_master as TPL


def _header_and_rows(csv_bytes: bytes) -> tuple[list[str], list[list[str]]]:
    text = csv_bytes.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    return rows[0], rows[1:]


# ───────────────────── #27 / #30 サポート提供 ─────────────────────

def test_column_lookup_ignores_surrounding_spaces():
    """page15 は取込時に header を strip する。空白を落とした名前でも照合できること。"""
    assert "サポート提供" in TPL.VALID_COLUMNS
    assert "直送アイテム" in TPL.VALID_COLUMNS


def test_support_column_always_emitted_with_default_t():
    """入力に「サポート提供」が無くても、CSV には必ず出て既定値 T が入る（#30）。"""
    header, rows = _header_and_rows(
        TPL.build_nst_master_csv([{"型番": "A-1", "アイテム名": "テスト品"}],
                                 ["型番", "アイテム名"])
    )
    assert "サポート提供 " in header, header      # 出力はテンプレ原文（末尾空白つき）
    assert rows[0][header.index("サポート提供 ")] == "T"


def test_support_column_keeps_explicit_value():
    """明示的に値が入っていれば既定値で上書きしない。"""
    header, rows = _header_and_rows(
        TPL.build_nst_master_csv([{"型番": "A-1", "サポート提供": "F"}],
                                 ["型番", "サポート提供"])
    )
    assert rows[0][header.index("サポート提供 ")] == "F"


def test_stripped_column_name_is_accepted_and_written_back_as_template_original():
    """strip 済みの名前で渡しても弾かれず、ヘッダはテンプレ原文で出る。"""
    header, rows = _header_and_rows(
        TPL.build_nst_master_csv([{"型番": "A-1", "直送アイテム": "T"}],
                                 ["型番", "直送アイテム"])
    )
    assert " 直送アイテム　" in header, header
    assert rows[0][header.index(" 直送アイテム　")] == "T"


def test_duplicate_after_normalization_is_not_emitted_twice():
    """原文名と strip 名を両方渡しても列は 1 本だけ。"""
    header, _ = _header_and_rows(
        TPL.build_nst_master_csv([{"型番": "A-1"}], ["型番", "サポート提供", "サポート提供 "])
    )
    assert header.count("サポート提供 ") == 1


def test_unknown_column_still_rejected():
    """正規化してもテンプレに無い列は従来どおり弾く（黙って通さない）。"""
    try:
        TPL.build_nst_master_csv([{"型番": "A-1"}], ["存在しない列"])
    except ValueError as e:
        assert "存在しない列" in str(e)
    else:
        raise AssertionError("非テンプレ列が素通りした")


# ───────────────────── 斑马 报关列（v2 · 2026-10-03） ─────────────────────

def _bm(row: dict) -> dict:
    return dict(zip(JBM.BM_HEADER, JBM.nst_to_bm_row(row)))


_PKG_ROW = {
    "JANコード": "4901234567890", "アイテム名": "テスト化粧水 200ml", "商品原価": "500",
    "パッケージ重量(g)": "230", "パッケージ奥行(cm)": "5", "パッケージ幅(cm)": "6",
    "パッケージ高さ(cm)": "18",
    "商品重量(g)": "999", "商品奥行(cm)": "99", "商品幅(cm)": "98", "商品高さ(cm)": "97",
    "_hs": "3304990000", "_name_en": "Test Lotion 200ml",
}


def test_bm_customs_columns_follow_spec():
    """中文名称=アイテム名 / 英文名称=_name_en / 报关重量=パッケージ重量 / 海关编码=_hs。"""
    b = _bm(_PKG_ROW)
    assert b["中文名称"] == "テスト化粧水 200ml"
    assert b["英文名称"] == "Test Lotion 200ml"
    assert b["报关重量(g)"] == "230"
    assert b["海关编码"] == "3304990000"


def test_bm_dimensions_use_package_only():
    """重量・长宽高はパッケージ。商品本体の値があっても使わない。"""
    b = _bm(_PKG_ROW)
    assert (b["重量(g)"], b["长(cm)"], b["宽(cm)"], b["高(cm)"]) == ("230", "5", "6", "18")


def test_bm_no_package_values_means_blank_not_product_fallback():
    """パッケージが空なら空。商品寸法へフォールバックしない。"""
    row = {k: v for k, v in _PKG_ROW.items() if not k.startswith("パッケージ")}
    b = _bm(row)
    assert (b["重量(g)"], b["报关重量(g)"], b["长(cm)"], b["宽(cm)"], b["高(cm)"]) == ("",) * 5


def test_bm_missing_name_en_and_hs_are_blank():
    """_name_en / _hs が無ければ空（日文名からの転写で埋めない）。"""
    b = _bm({"JANコード": "4901234567890", "アイテム名": "テスト品"})
    assert b["英文名称"] == ""
    assert b["海关编码"] == ""


def test_bm_name_en_capped_at_76():
    b = _bm({**_PKG_ROW, "_name_en": "A" * 100})
    assert len(b["英文名称"]) == 76


def test_bm_xlsx_layout_unchanged():
    """46 列ヘッダ・row1 結合セル・シート名は変えない。データは row3 から。"""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(JBM.build_bm_xlsx([_PKG_ROW])))
    assert wb.sheetnames == ["数据"]
    ws = wb["数据"]
    assert [c.value for c in ws[2]] == JBM.BM_HEADER and len(JBM.BM_HEADER) == 46
    assert sorted(str(r) for r in ws.merged_cells.ranges) == sorted(JBM.BM_MERGES)
    data = dict(zip(JBM.BM_HEADER, [c.value for c in ws[3]]))
    assert data["海关编码"] == "3304990000" and data["英文名称"] == "Test Lotion 200ml"


def test_jd_output_removed():
    assert not hasattr(JBM, "build_jd_xlsx") and not hasattr(JBM, "nst_to_jd_row")
