"""商品登録 v2 入力テンプレ: 生成物の構造と read_upload の判定。"""
from __future__ import annotations

import io

import openpyxl
import pytest

from data_warehouse.templates import item_entry_form as F

try:
    from shared.nst_choices import Choices
except ImportError:  # A が並行で書いている間の保険（インタフェース通りの形）
    from dataclasses import dataclass

    @dataclass(frozen=True)
    class Choices:
        suppliers: tuple
        middle_by_large: dict
        owners: tuple = ("037 米澤和敏", "043 徐越", "005 川崎里子", "079 隋艶偉")
        warnings: tuple = ()

        @property
        def large(self):
            return tuple(self.middle_by_large)


SUPPLIERS = ("0001 ベストアンサ株式会社", "0068 SHANGHAI SHENGTUO INTERNATIONAL TRADE,Co.Ltd") + tuple(
    f"{i:04d} テスト商事株式会社" for i in range(100, 699))   # 601 件・カンマ入りあり
CH = Choices(
    suppliers=SUPPLIERS,
    middle_by_large={"アウトドア・防災": ("アウトドア用品", "防災用品"),
                     "ASEAN": ("日用消耗品", "その他（生活雑貨）"),
                     "USA": ("日用消耗品",), "食品": ("加工食品",)},
    owners=("037 米澤和敏", "043 徐越", "005 川崎里子", "079 隋艶偉"),
)


def _cd(body: str) -> str:
    """チェックディジット（左から奇数位×1・偶数位×3 で書いた独立実装）。"""
    pad = body.rjust(12, "0")
    s = sum(int(c) * (1 if k % 2 == 0 else 3) for k, c in enumerate(pad))
    return body + str((10 - s % 10) % 10)


J1, J2, J3 = _cd("490123456789"), _cd("454929200001"), _cd("4901234")   # J3 は 8 桁


def test_check_digit():
    assert J1 == "4901234567894" and J3 == "49012347"
    assert F.jan_check_ok(J1) and F.jan_check_ok(J3)
    assert not F.jan_check_ok("4901234567890")


def test_template_structure():
    wb = openpyxl.load_workbook(io.BytesIO(F.build_template(CH)))
    assert wb.sheetnames == ["商品登録", "記入例", "_選択肢"]
    assert wb["_選択肢"].sheet_state == "hidden"
    ws = wb["商品登録"]
    assert [c.value for c in ws[1]] == F.INPUT_COLUMNS
    assert ws.freeze_panes == "B2"
    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 4
    for dv in dvs:
        assert dv.type == "list" and dv.showErrorMessage is True and not dv.showDropDown
        assert len(dv.formula1) < 255
        assert str(dv.sqref).endswith("1001")
    by_col = {str(dv.sqref)[0]: dv.formula1 for dv in dvs}
    assert by_col["C"] == "'_選択肢'!$A$2:$A$602"      # 仕入先 601 件は範囲参照
    assert by_col["F"].startswith("INDIRECT(VLOOKUP($E2,")
    # 定義名 MID_xx → 各大分類の中分類列
    ch = wb["_選択肢"]
    names = dict(wb.defined_names.items())
    for r in range(2, 2 + len(CH.large)):
        lg, alias = ch.cell(r, 3).value, ch.cell(r, 4).value
        ref = names[alias].attr_text.split("!")[1].replace("$", "")
        col = ref.split(":")[0].rstrip("0123456789")
        vals = tuple(ch[f"{col}{k}"].value for k in range(2, 2 + len(CH.middle_by_large[lg])))
        assert ref.endswith(str(1 + len(vals)))
        assert ch[f"{col}1"].value == lg and vals == CH.middle_by_large[lg]
    assert wb["記入例"].cell(2, 1).value == "4901234567894"


def _upload(rows: list[list], header=None) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(F.build_template(CH)))
    ws = wb["商品登録"]
    if header:
        for j, h in enumerate(header, start=1):
            ws.cell(1, j, h)
    for i, r in enumerate(rows, start=2):
        for j, v in enumerate(r, start=1):
            ws.cell(i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


OK = [J1, 1250, SUPPLIERS[1], "043 徐越", "ASEAN", "日用消耗品", 48, 12, 6, 23, 17, 250]


def _row(**kw):
    r = list(OK)
    for k, v in kw.items():
        r[F.INPUT_COLUMNS.index(k)] = v
    return r


def test_read_upload_ok_and_errors():
    rows_in = [
        OK,                                                                   # 2 ok
        _row(**{"JANコード": float(J2), "商品原価": " 980.5 ", "パッケージ幅(cm)": None,
                "パッケージ重量(g)": "約70"}),                                  # 3 ok + warn
        [None] * 12,                                                          # 4 空行 → 無視
        _row(**{"JANコード": J1}),                                            # 5 重複
        _row(**{"JANコード": "4901234567890"}),                               # 6 CD 不一致
        _row(**{"JANコード": "12345"}),                                       # 7 桁数
        _row(**{"JANコード": J3, "商品原価": 0}),                             # 8 原価
        _row(**{"JANコード": _cd("490000000001"), "仕入先1：仕入先": "9999 無い"}),   # 9
        _row(**{"JANコード": _cd("490000000002"), "商品担当者": "007 安部龍也"}),     # 10
        _row(**{"JANコード": _cd("490000000003"), "大分類": "中国"}),          # 11
        _row(**{"JANコード": _cd("490000000004"), "中分類": "加工食品"}),      # 12 組合せ
        _row(**{"JANコード": _cd("490000000005"), "カートン入数": 1.5}),       # 13
        _row(**{"JANコード": _cd("490000000006"), "発注ロット": None,
                "カートン入数": "２４"}),                                       # 14 ok + warn
    ]
    rows, issues = F.read_upload(_upload(rows_in), CH)
    assert [r["JANコード"] for r in rows] == [J1, J2, _cd("490000000006")]
    assert set(rows[0]) == set(F.INPUT_COLUMNS)
    assert all(isinstance(v, str) for r in rows for v in r.values())
    assert rows[0]["商品原価"] == "1250" and rows[0]["パッケージ高さ(cm)"] == "6"
    assert rows[1]["商品原価"] == "980.5"
    assert rows[1]["パッケージ幅(cm)"] == "" and rows[1]["パッケージ重量(g)"] == ""
    assert rows[2]["カートン入数"] == "24" and rows[2]["発注ロット"] == ""

    err_rows = sorted({i.row for i in issues if i.level == "error"})
    assert err_rows == [5, 6, 7, 8, 9, 10, 11, 12, 13]
    msg = {i.row: i.message for i in issues if i.level == "error"}
    assert "重複" in msg[5] and "2 行目" in msg[5]
    assert "チェックディジット" in msg[6]
    assert "8 桁 / 13 桁" in msg[7]
    assert "商品原価" in msg[8]
    assert "仕入先" in msg[9] and "商品担当者" in msg[10] and "大分類が選択肢" in msg[11]
    assert "中分類「加工食品」" in msg[12] and "カートン入数" in msg[13]
    warn_rows = sorted({i.row for i in issues if i.level == "warn"})
    assert warn_rows == [3, 14]
    w3 = [i.message for i in issues if i.row == 3]
    assert any("パッケージ 3 辺" in m for m in w3) and any("パッケージ重量" in m for m in w3)


def test_columns_reordered_and_header_missing():
    order = list(reversed(F.INPUT_COLUMNS))
    rows, issues = F.read_upload(_upload([list(reversed(OK))], header=order), CH)
    assert issues == [] and rows[0]["JANコード"] == J1 and rows[0]["仕入先1：仕入先"] == SUPPLIERS[1]

    bad = list(F.INPUT_COLUMNS)
    bad[1] = "原価"
    rows, issues = F.read_upload(_upload([OK], header=bad), CH)
    assert rows == [] and issues[0].level == "error" and "商品原価" in issues[0].message


def test_wrong_sheet():
    wb = openpyxl.Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    rows, issues = F.read_upload(buf.getvalue(), CH)
    assert rows == [] and "商品登録" in issues[0].message


def test_filename():
    assert F.template_filename().endswith(".xlsx")
