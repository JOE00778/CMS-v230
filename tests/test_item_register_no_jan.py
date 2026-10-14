"""No-JAN registration: independent item codes survive mixed imports and exports."""
import csv
import io
import pytest
from unittest.mock import MagicMock

import openpyxl
from data_warehouse.templates import item_entry_form as F, jd_bm_item_master as BM, nst_item_master as NST
from shared import item_register as R, nst_choices as NC
from shared.hs_classify import Customs
from shared.jan_web import JanMeta

SKUS = ['ETERNAL-BONSAI-' + c for c in ('REDN', 'NV', 'GR', 'WH', 'BK', 'LGR')]
JAN = '4936201110209'
CH = NC.Choices(suppliers=('0001 TEST',), middle_by_large={'生活雑貨': ('日用品雑貨',)})


def upload(records):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = F.SHEET_INPUT
    columns = [*F.INPUT_COLUMNS, F.COL_NAME_MANUAL, '型番', 'メーカー名（任意）']
    ws.append(columns)
    for values in records:
        r = {'JANコード': '', '商品原価': '100', '仕入先1：仕入先': '0001 TEST',
             '商品担当者': '043 徐越', '大分類': '生活雑貨', '中分類': '日用品雑貨',
             'カートン入数': '1', '発注ロット': '1', F.COL_NAME_MANUAL: '盆栽',
             'メーカー名（任意）': '手入力メーカー', **values}
        ws.append([r.get(c, '') for c in columns])
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def test_upload_accepts_six_no_jan_and_rejects_duplicate_codes_and_bad_jan():
    records = [{'型番': s} for s in SKUS] + [{'JANコード': JAN}, {'型番': SKUS[0]},
              {'型番': 'BAD', 'JANコード': '123'}, {'型番': 'NO-NAME', F.COL_NAME_MANUAL: ''}]
    rows, issues = F.read_upload(upload(records), CH)
    assert len(rows) == 7
    assert [r['型番'] for r in rows] == [*SKUS, JAN]
    assert {i.row for i in issues if i.level == 'error'} == {9, 10, 11}


@pytest.mark.parametrize("jan_code", [JAN, "CREAM-SKU"])
def test_no_jan_mixed_pipeline_preserves_identity_and_skips_barcode_queries(jan_code):
    rows, _ = F.read_upload(upload([{'型番': s} for s in SKUS] + [{'JANコード': JAN, '型番': jan_code}]), CH)
    meta, weight = MagicMock(return_value=JanMeta(JAN, 'クリーム', 'AUTO', '', 'ok')), MagicMock(return_value=None)
    nst, extra, issues, stats = R.enrich(rows, choices=CH, prefix_makers={}, fetch_meta=meta,
        fetch_weight=weight, classify_customs=lambda *a, **k: Customs(None, None, None, 'none'),
        want_weight=True, interval=0)
    assert len(nst) == 7 and set(extra) == {*SKUS, jan_code}
    assert meta.call_count == weight.call_count == 1
    assert meta.call_args.args == weight.call_args.args == (JAN,)
    assert all(r['メーカー名'] == '手入力メーカー' for r in nst)
    frame = R.to_frame(nst, extra)
    for i, r in enumerate(frame): r['アイテム名'] = f'編集 {i}'
    out, bm, sources, _ = R.from_frame(frame, nst, extra)
    assert [r['型番'] for r in out] == [*SKUS, jan_code]
    assert [r['アイテム名'] for r in out] == [f'編集 {i}' for i in range(7)]
    data = list(csv.DictReader(io.StringIO(NST.build_nst_master_csv(out, R.csv_field_columns(out), id_label='型番').decode('utf-8-sig'))))
    assert [r['JANコード'] for r in data] == [''] * 6 + [JAN]
    mapped = [dict(zip(BM.BM_HEADER, BM.nst_to_bm_row(r))) for r in bm]
    assert [r['SKU'] for r in mapped] == [*SKUS, jan_code]
    assert [r['SPU'] for r in mapped] == [*SKUS, jan_code]
    calls = []
    def ai(items):
        calls.extend(items); return {code: 'mock failure' for code, _, _ in items}
    R.cross_check_hs(out, extra, stats, ai)
    assert [x[0] for x in calls] == [*SKUS, jan_code]
    conn = MagicMock()
    assert R.write_register_log(conn, out, None, sources) == (7, '')
    sql, params = conn.executemany.call_args.args
    assert 'item_code' in sql
    assert [p[0] for p in params] == [*SKUS, jan_code]
    assert [p[1] for p in params] == [''] * 6 + [JAN]


def test_existing_item_codes_query_uses_itemid_and_pg_fallback():
    calls = []
    def query(sql, **kw): calls.append(sql); return [{'itemid': SKUS[0]}]
    found, warnings = NC.existing_item_codes(SKUS, suiteql=query)
    assert found == {SKUS[0]} and not warnings and 'itemid IN' in calls[0]
    conn = MagicMock(); conn.execute.return_value.fetchall.return_value = [{'item_code': SKUS[1]}]
    def failed(*a, **kw): raise RuntimeError('offline')
    found, warnings = NC.existing_item_codes(SKUS, suiteql=failed, conn=conn)
    assert found == {SKUS[1]} and warnings


def test_page15_no_jan_generation_and_existing_code_exclusion():
    import pathlib
    import zipfile
    from unittest.mock import patch
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    from shared import jan_web, hs_ai
    st.cache_data.clear()
    up = MagicMock(); up.getvalue.return_value = upload([{'型番': s} for s in SKUS])
    conn = MagicMock(); conn.__enter__.return_value = conn
    with patch('shared.db.get_connection', return_value=conn), \
         patch('streamlit.file_uploader', return_value=up), \
         patch.object(NC, 'load_choices', return_value=CH), \
         patch.object(NC, 'load_prefix_makers', return_value={}), \
         patch.object(NC, 'existing_jans', return_value=(set(), [])), \
         patch.object(NC, 'existing_item_codes', return_value=({SKUS[0]}, [])), \
         patch.object(hs_ai, 'is_configured', return_value=False), \
         patch.object(jan_web, 'fetch_meta') as meta:
        page = pathlib.Path(__file__).resolve().parents[1] / 'pages' / '15_📝_商品登录.py'
        at = AppTest.from_file(str(page), default_timeout=30)
        at.session_state['__auth_ok'] = True
        at.run(); at.button(key='page15_run').click().run()
        assert not at.exception, [e.value for e in at.exception]
        assert [m.value for m in at.metric][:2] == ['5', '1']
        checklist = next(d.value for d in at.dataframe if '待手入力' in d.value.columns)
        assert checklist['型番'].tolist() == SKUS[1:]
        assert all('メーカー名' not in text for text in checklist['待手入力'])
        assert all(F.COL_PKG_WEIGHT not in text for text in checklist['待手入力'])
        editor = next(d.value for d in at.dataframe if R.COL_PKG_WEIGHT in d.value.columns)
        assert all(editor[c].eq('1').all() for c in R._WEIGHT_FIELDS)
        at.button(key='page15_gen').click().run()
        assert not at.exception, [e.value for e in at.exception]
        assert at.session_state['page15_zip_n'] == 5
        assert not meta.called
        with zipfile.ZipFile(io.BytesIO(at.session_state['page15_zip_bytes'])) as z:
            nst = next(n for n in z.namelist() if n.endswith('.csv'))
            rows = list(csv.DictReader(io.StringIO(z.read(nst).decode('utf-8-sig'))))
            assert [r['型番'] for r in rows] == SKUS[1:]
            bm_name = next(n for n in z.namelist() if n.endswith('.xlsx'))
            ws = openpyxl.load_workbook(io.BytesIO(z.read(bm_name)))['数据']
            assert [ws.cell(i, 13).value for i in range(3, 8)] == SKUS[1:]


def test_manual_input_checklist_reports_only_missing_no_jan_fields():
    fields = ['メーカー名', *F.COL_PKG_DIMS, F.COL_PKG_WEIGHT, 'HS', '通関英文名']
    partial = {'型番': SKUS[0], 'JANコード': '', 'アイテム名': '盆栽',
               'メーカー名': '確認済メーカー', F.COL_PKG_DIMS[0]: '10', 'HS': '999999'}
    complete = {'型番': SKUS[1], 'JANコード': '', 'アイテム名': '盆栽',
                **{f: 'filled' for f in fields}}
    regular = {'型番': JAN, 'JANコード': JAN}
    needs = R.manual_input_needs([partial, complete, regular])
    assert len(needs) == 1 and needs[0]['型番'] == SKUS[0]
    text = needs[0]['待手入力']
    assert 'メーカー名' not in text and F.COL_PKG_DIMS[0] not in text and 'HS' not in text
    assert text == '通関英文名'
    assert '確認・修正' in needs[0]['入力場所']
    assert R.manual_input_needs([complete, regular]) == []


def test_manual_input_checklist_points_identity_and_name_to_template():
    needs = R.manual_input_needs([{'JANコード': '', '型番': '', 'アイテム名': ''}])
    assert '型番' in needs[0]['待手入力'] and 'アイテム名' in needs[0]['待手入力']
    assert 'テンプレ' in needs[0]['入力場所']


def test_bonsai_models_in_jan_column_are_converted_and_exported_as_codes():
    rows, issues = F.read_upload(upload([{'JANコード': s} for s in SKUS]), CH)
    assert len(rows) == 6 and not any(i.level == 'error' for i in issues)
    assert [r['型番'] for r in rows] == SKUS
    assert all(r['JANコード'] == '' for r in rows)
    meta, weight = MagicMock(), MagicMock()
    nst, extra, _, _ = R.enrich(rows, choices=CH, prefix_makers={}, fetch_meta=meta,
        fetch_weight=weight, classify_customs=lambda *a, **k: Customs(None, None, None, 'none'),
        want_weight=True, interval=0)
    assert not meta.called and not weight.called
    out, bm, _, _ = R.from_frame(R.to_frame(nst, extra), nst, extra)
    csv_rows = list(csv.DictReader(io.StringIO(NST.build_nst_master_csv(out,
        R.csv_field_columns(out), id_label='型番').decode('utf-8-sig'))))
    assert [r['型番'] for r in csv_rows] == SKUS
    assert all(not r.get('JANコード') for r in csv_rows)
    assert [dict(zip(BM.BM_HEADER, BM.nst_to_bm_row(r)))['SKU'] for r in bm] == SKUS


def test_bonsai_jan_alias_rejects_conflicting_codes_and_other_invalid_barcodes():
    records = [{'JANコード': SKUS[0], '型番': 'OTHER'}, {'JANコード': 'ETERNAL-BONSAI-UNKNOWN'},
               {'JANコード': 'OTHER-SKU'}, {'JANコード': '4936201110200'},
               {'JANコード': SKUS[1], '型番': SKUS[1]}, {'JANコード': '', '型番': SKUS[1]}]
    rows, issues = F.read_upload(upload(records), CH)
    assert [r['型番'] for r in rows] == [SKUS[1]]
    assert {i.row for i in issues if i.level == 'error'} == {2, 3, 4, 5, 7}
    assert any('一致' in i.message for i in issues if i.row == 2)
