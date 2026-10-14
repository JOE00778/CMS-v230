"""商品登録 v2 の入力テンプレ（Excel）生成とアップロード解析。DB 非依存。

旧 INPUT_COLUMNS の 12 列を維持し、原本の手入力・条件項目を追加。列名は NST テンプレ原本
（nst_item_master_template.xlsx · NetSuiteマスタ登録 row2）の原文そのまま。
固定値列・アイテム名・メーカー名・税率などは呼び出し側が NST 行を組むときに足す。

シート構成:
  商品登録  … row1=見出し（原文列名・説明はセルのコメント）/ row2〜=入力。原本 NetSuiteマスタ登録 シートの row5〜も読める
  記入例    … 見本 1 行（読まない）
  _選択肢   … hidden。仕入先 / 商品担当者 / 大分類→定義名 対応表 / 大分類ごとの中分類列

中分類の連動（2026-10-03 実測して決めた）:
  openpyxl 3.1.5 は中黒入りの定義名（アウトドア・防災）も書けて読み戻せる。ただし
  Excel/LibreOffice で開いての確認は Mac に無いため不可。大分類は SuiteQL から動的に来るので、
  将来 空白・記号・セル番地に見える名前（例 "USA" は列名としても有効）が来ても壊れないよう、
  定義名は安全な別名 MID_01.. にし、中分類の入力規則は
  INDIRECT(VLOOKUP(大分類, 対応表, 2, 0)) で引く。"L01" は Excel が L1 セルと解釈しうるので避けた。
  連動はあくまで入力補助。正しさは read_upload の (大分類, 中分類) 組合せ検証で担保する。
"""
from __future__ import annotations

import datetime
import io
import math
import re
import unicodedata
from dataclasses import dataclass

from data_warehouse.templates.nst_item_master import (
    NST_MASTER_COLUMNS, NST_FIELD_HELP, NST_AUTOMATIC_COLUMNS, NST_REGISTRATION_FIXED, _norm,
)

COL_JAN = "JANコード"
COL_COST = "商品原価"
COL_SUPPLIER = "仕入先1：仕入先"
COL_OWNER = "商品担当者"
COL_LARGE = "大分類"
COL_MIDDLE = "中分類"
COL_CARTON = "カートン入数"
COL_LOT = "発注ロット"
COL_PKG_DIMS = ("パッケージ高さ(cm)", "パッケージ幅(cm)", "パッケージ奥行(cm)")
COL_PKG_WEIGHT = "パッケージ重量(g)"
COL_CARTON_DIMS = ("カートン高さ(cm)", "カートン幅(cm)", "カートン奥行(cm)")

INPUT_COLUMNS: list[str] = [COL_JAN, COL_COST, COL_SUPPLIER, COL_OWNER, COL_LARGE, COL_MIDDLE,
                            COL_CARTON, COL_LOT, *COL_PKG_DIMS, COL_PKG_WEIGHT]
# 2026-10-03: jancode.xyz に商品名が無い JAN を登録できるよう任意列を足した。NST の列ではないので
# INPUT_COLUMNS には入れない（無い旧テンプレもそのまま読める）。rows には常にキーとして入る（空可）
COL_NAME_MANUAL = "アイテム名（任意）"
COL_ITEM_CODE = "型番"
COL_MAKER_MANUAL = "メーカー名（任意）"
# Boss 2026-10-13: この 6 型番だけは JAN 欄への入力を互換受付。実 JAN は空。
JAN_MODEL_ALIASES = frozenset("ETERNAL-BONSAI-" + c for c in ("REDN", "NV", "GR", "WH", "BK", "LGR"))
TEMPLATE_COLUMNS: list[str] = [COL_JAN, COL_NAME_MANUAL, *INPUT_COLUMNS[1:], COL_ITEM_CODE, COL_MAKER_MANUAL]

# 旧入力列は必須ヘッダとして保持。新列は末尾に追加、旧テンプレも読める。
ADDITIONAL_INPUT_COLUMNS = [c for _, c in NST_MASTER_COLUMNS
                           if c not in (*INPUT_COLUMNS, COL_ITEM_CODE, "アイテム名", "メーカー名",
                                        *NST_AUTOMATIC_COLUMNS, *NST_REGISTRATION_FIXED)]
TEMPLATE_COLUMNS += ADDITIONAL_INPUT_COLUMNS
NUMERIC_COLUMNS = tuple(c for _, c in NST_MASTER_COLUMNS
                        if "(cm)" in c or "(g)" in c) + (COL_CARTON, COL_COST,
    "コンテナ運送費", "開発費単価", "不良率自社負担単価", "不良率(PO)", "関税(AT)",
    "関税率", "ライセンス料(AT)", "ライセンス率", "その他費用", "参考販売価格(税抜)", "直送送料")
FIELD_CHOICES = {
    "商品ランク": ("NEW", "Aランク", "Bランク", "Cランク", "取扱中止"),
    "保証期間": ("保証なし", "3ヶ月", "6ヶ月", "1年", "2年", "3年"),
    "不良対応方法（社内）": ("返送", "お客様処分（写真・動画必要）", "お客様処分（報告のみ）"),
    "不良対応方法（仕入先）": ("仕入先返品", "処分（不良箇所写真・動画）", "処分（仕入先指定写真・動画）", "処分（対応なし）", "保管", "パーツ交換"),
    "税率": ("★",), " 直送アイテム　": ("T",), "特注アイテム": ("T",),
    "商品サイズ": ("T", "F"), "段ボール梱包": ("T", "F"),
}

SHEET_INPUT = "商品登録"
SHEET_EXAMPLE = "記入例"
SHEET_CHOICES = "_選択肢"
FIRST_DATA_ROW = 2
LAST_DATA_ROW = 1001

_REQUIRED = {name for mark, name in NST_MASTER_COLUMNS if mark == "必須"} - {COL_JAN, COL_ITEM_CODE}

# 見出しコメント（原本 row3 を短くしたもの）
_HELP = {
    COL_JAN: "JANコード（8 桁 or 13 桁）。JAN の無い商品は空欄、型番欄に商品コードを記入。ETERNAL-BONSAI の既定 6 型番はこの JAN 欄に直接入力可（型番として処理）",
    COL_ITEM_CODE: "商品を一意に識別する型番 / SKU（最大 250 字）。空なら JAN を使用。JAN が無い場合は必須",
    COL_MAKER_MANUAL: "確認済みメーカー名。記入値を優先。不明なら空欄のまま（任意・最大 60 字）",
    COL_NAME_MANUAL: "通常は空のまま（jancode.xyz の商品名が自動で入ります）。"
                     "JAN が無い商品・jancode に商品名が無い商品は必ず記入。記入した場合はこちらが優先。最大 60 字",
    COL_COST: "商品の原価（日本円・仕入価格）。「定義原価」ではありません",
    COL_SUPPLIER: "選択肢以外入力禁止。無い場合は「NetSuite【仕入先】マスタ登録依頼書」で登録依頼",
    COL_OWNER: "このアイテムの担当者を選択",
    COL_LARGE: "大分類を選択",
    COL_MIDDLE: "大分類を先に選ぶと、その中分類だけが出ます",
    COL_CARTON: "カートン入数（正の数。原本は整数または小数）",
    COL_LOT: "発注時のロット数・条件（原本例：1色1000、最大 60 字）",
    COL_PKG_DIMS[0]: "外装パッケージの高さ（整数または小数）。空欄は網調べ後も不明なら 1 cm の仮値（非実測）",
    COL_PKG_DIMS[1]: "外装パッケージの幅（整数または小数）。不明なら 1 cm の仮値（非実測）",
    COL_PKG_DIMS[2]: "外装パッケージの奥行（整数または小数）。不明なら 1 cm の仮値（非実測）",
    COL_PKG_WEIGHT: "外装込みの重量 g（整数または小数）。不明なら 1 g の仮値（非実測）",
}


@dataclass
class Issue:
    row: int        # Excel 上の行番号（見出し不備などファイル全体の問題は 0）
    jan: str       # 互換フィールド名。画面では型番 / JAN の識別子を表示
    level: str      # "error" | "warn"
    message: str


def template_filename() -> str:
    return f"商品登録_入力テンプレ_{datetime.date.today():%Y%m%d}.xlsx"


# ───────────────────────── テンプレ生成 ─────────────────────────

def _quoted(sheet: str) -> str:
    return "'" + sheet.replace("'", "''") + "'"


def _list_dv(formula1: str, col_letter: str, title: str):
    from openpyxl.worksheet.datavalidation import DataValidation
    # showErrorMessage は openpyxl 既定 False（選択肢外も通る）なので必ず True。showDropDown は触らない
    dv = DataValidation(type="list", formula1=formula1, allow_blank=True,
                        showErrorMessage=True, errorStyle="stop",
                        errorTitle=title, error="選択肢から選んでください")
    dv.add(f"{col_letter}{FIRST_DATA_ROW}:{col_letter}{LAST_DATA_ROW}")
    return dv


def _write_header(ws) -> None:
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    req_fill = PatternFill("solid", fgColor="F4B084")
    opt_fill = PatternFill("solid", fgColor="D9D9D9")
    for i, name in enumerate(TEMPLATE_COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True)
        c.fill = req_fill if name in _REQUIRED else opt_fill
        mark = next((m for m, col in NST_MASTER_COLUMNS if col == name), None)
        c.comment = Comment((f"【{mark}】" if mark else "") + _HELP.get(name, NST_FIELD_HELP.get(name, "")), "CMS")
        ws.column_dimensions[c.column_letter].width = 34 if name in (COL_SUPPLIER, COL_NAME_MANUAL) else 16


def build_template(choices) -> bytes:
    """choices は shared.nst_choices.Choices（suppliers / owners / middle_by_large / large）。"""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_INPUT
    ex = wb.create_sheet(SHEET_EXAMPLE)
    ch = wb.create_sheet(SHEET_CHOICES)
    ch.sheet_state = "hidden"
    q = _quoted(SHEET_CHOICES)

    # _選択肢: A=仕入先 B=商品担当者 C=大分類 D=定義名 E〜=大分類ごとの中分類
    large = list(choices.large)
    for col, title, values in ((1, "仕入先", choices.suppliers), (2, "商品担当者", choices.owners),
                               (3, "大分類", large),
                               (4, "定義名", [f"MID_{i:02d}" for i in range(1, len(large) + 1)])):
        ch.cell(row=1, column=col, value=title)
        for r, v in enumerate(values, start=2):
            ch.cell(row=r, column=col, value=v)
    for i, lg in enumerate(large, start=1):
        col = 4 + i
        letter = get_column_letter(col)
        middles = choices.middle_by_large[lg]
        ch.cell(row=1, column=col, value=lg)
        for r, v in enumerate(middles, start=2):
            ch.cell(row=r, column=col, value=v)
        end = max(len(middles), 1) + 1
        wb.defined_names[f"MID_{i:02d}"] = DefinedName(
            f"MID_{i:02d}", attr_text=f"{q}!${letter}$2:${letter}${end}")

    def rng(letter: str, n: int) -> str:
        return f"{q}!${letter}$2:${letter}${max(n, 1) + 1}"

    _write_header(ws)
    ws.freeze_panes = "B2"
    letter_of = {name: get_column_letter(i) for i, name in enumerate(TEMPLATE_COLUMNS, start=1)}
    for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):   # JAN は文字列で入れてもらう（指数表記・桁落ち防止）
        ws[f"{letter_of[COL_JAN]}{r}"].number_format = "@"
    large_cell = f"${letter_of[COL_LARGE]}{FIRST_DATA_ROW}"
    for dv in (
        _list_dv(rng("A", len(choices.suppliers)), letter_of[COL_SUPPLIER], COL_SUPPLIER),
        _list_dv(rng("B", len(choices.owners)), letter_of[COL_OWNER], COL_OWNER),
        _list_dv(rng("C", len(large)), letter_of[COL_LARGE], COL_LARGE),
        _list_dv(f"INDIRECT(VLOOKUP({large_cell},{q}!$C$2:$D${max(len(large), 1) + 1},2,0))",
                 letter_of[COL_MIDDLE], COL_MIDDLE),
    ):
        ws.add_data_validation(dv)

    for col, values in FIELD_CHOICES.items():
        if col in letter_of:
            ws.add_data_validation(_list_dv('"' + ",".join(values) + '"', letter_of[col], col))

    # 記入例
    _write_header(ex)
    lg0 = large[0] if large else ""
    sample = {COL_JAN: "4901234567894", COL_NAME_MANUAL: "", COL_COST: 1250,
              COL_SUPPLIER: choices.suppliers[0] if choices.suppliers else "",
              COL_OWNER: choices.owners[0] if choices.owners else "",
              COL_LARGE: lg0, COL_MIDDLE: (choices.middle_by_large.get(lg0) or ("",))[0],
              COL_CARTON: 48, COL_LOT: 12, COL_PKG_DIMS[0]: 6, COL_PKG_DIMS[1]: 23,
              COL_PKG_DIMS[2]: 17, COL_PKG_WEIGHT: 250}
    for i, name in enumerate(TEMPLATE_COLUMNS, start=1):
        ex.cell(row=2, column=i, value=sample.get(name, ""))
    ex.cell(row=4, column=1, value="※ このシートは読み込まれません。「商品登録」シートに記入してください")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ───────────────────────── アップロード解析 ─────────────────────────

def _s(v) -> str:
    """セル値を文字列に。Excel 由来の 4901234567890.0 を潰し、前後空白を落とす。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip().strip("　").strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def _num(s: str) -> float | None:
    try:
        n = float(s)
        return n if math.isfinite(n) else None
    except ValueError:
        return None


def _nfkc(s: str) -> str:
    """数値系の列だけ全角数字・全角ピリオドを半角に。"""
    return _s(unicodedata.normalize("NFKC", s))


def jan_check_ok(jan: str) -> bool:
    """GTIN（JAN-8 / JAN-13）のチェックディジット。右端の 1 つ手前から 3,1,3,1… の重み。"""
    d = [int(c) for c in jan]
    total = sum(x * (3 if i % 2 == 0 else 1) for i, x in enumerate(reversed(d[:-1])))
    return (10 - total % 10) % 10 == d[-1]


def _open(file):
    from openpyxl import load_workbook
    if isinstance(file, (bytes, bytearray)):
        file = io.BytesIO(file)
    return load_workbook(file, data_only=True, read_only=True)


def read_upload(file, choices) -> tuple[list[dict], list[Issue]]:
    """「商品登録」シートを読む。正常行（warn 込み）だけ rows で返し、error 行は issues にだけ入れる。

    件数: ok = len(rows) / error 行数 = len({i.row for i in issues if i.level == "error" and i.row})
          / warn 行数 = len({i.row for i in issues if i.level == "warn"})
    """
    wb = _open(file)
    sheet_name = SHEET_INPUT if SHEET_INPUT in wb.sheetnames else "NetSuiteマスタ登録"
    if sheet_name not in wb.sheetnames:
        return [], [Issue(0, "", "error", f"シート「{SHEET_INPUT}」がありません")]
    all_rows = list(wb[sheet_name].iter_rows(values_only=True))
    wb.close()

    # 見出し行: 先頭 10 行で JANコード を含む最初の行
    header_idx = next((i for i, r in enumerate(all_rows[:10])
                       if any(_norm(_s(v)) == COL_JAN for v in r)), None)
    if header_idx is None:
        return [], [Issue(0, "", "error", f"見出し「{COL_JAN}」が見つかりません")]
    pos = {}
    for j, v in enumerate(all_rows[header_idx]):
        pos.setdefault(_norm(_s(v)), j)
    missing = [c for c in INPUT_COLUMNS if _norm(c) not in pos]
    if missing:
        return [], [Issue(header_idx + 1, "", "error", "見出しが見つかりません: " + " / ".join(missing))]

    suppliers, owners = set(choices.suppliers), set(choices.owners)
    mbl = choices.middle_by_large
    rows: list[dict] = []
    issues: list[Issue] = []
    seen: dict[str, int] = {}
    seen_codes: dict[str, int] = {}

    data_idx = header_idx + (3 if sheet_name == "NetSuiteマスタ登録" else 1)
    for i, raw in enumerate(all_rows[data_idx:], start=data_idx + 1):
        rec = {c: _s(raw[pos[_norm(c)]]) if pos[_norm(c)] < len(raw) else "" for c in INPUT_COLUMNS}
        p_name = pos.get(_norm(COL_NAME_MANUAL), pos.get("アイテム名"))
        rec[COL_NAME_MANUAL] = _s(raw[p_name]) if p_name is not None and p_name < len(raw) else ""
        for col in (COL_ITEM_CODE, COL_MAKER_MANUAL):
            idx = pos.get(_norm(col), pos.get("メーカー名") if col == COL_MAKER_MANUAL else None)
            rec[col] = _s(raw[idx]) if idx is not None and idx < len(raw) else ""
        for col in ADDITIONAL_INPUT_COLUMNS:
            idx = pos.get(_norm(col))
            val = raw[idx] if idx is not None and idx < len(raw) else None
            rec[col] = val.strftime("%Y/%m/%d") if col == "発売日" and isinstance(val, datetime.date) else _s(val)
            if col in ("商品サイズ", "段ボール梱包") and rec[col] in ("True", "False"):
                rec[col] = "T" if rec[col] == "True" else "F"
        if not any(rec.get(c) for c in (COL_JAN, COL_ITEM_CODE, COL_NAME_MANUAL, COL_COST, COL_SUPPLIER)):
            continue
        for c in (COL_JAN, *NUMERIC_COLUMNS):
            rec[c] = _nfkc(rec[c])
        jan = rec[COL_JAN]
        errs: list[str] = []
        warns: list[str] = []
        from shared.nst_choices import FIXED
        for col in NST_REGISTRATION_FIXED:
            idx = pos.get(_norm(col))
            value = _s(raw[idx]) if idx is not None and idx < len(raw) else ""
            if value and value != FIXED[col]:
                errs.append(f"{col}={value} は現在の登録フロー対象外です（{FIXED[col]}専用）。固定値へ勝手に変更せず、別途確認してください")

        # NST 原本では JAN なしのとき型番を JAN 欄に記入。内部は実 JAN を分離する。
        if rec[COL_ITEM_CODE] and jan == rec[COL_ITEM_CODE] and not (jan.isdigit() and len(jan) in (8, 13)):
            rec[COL_JAN] = jan = ""
        if jan in JAN_MODEL_ALIASES:
            if rec[COL_ITEM_CODE] and rec[COL_ITEM_CODE] != jan:
                errs.append("JAN 欄の盆栽型番と型番欄の値が一致しません（同じ型番にしてください）")
            else:
                rec[COL_ITEM_CODE] = jan
                rec[COL_JAN] = ""
                jan = ""

        if jan and not (jan.isdigit() and len(jan) in (8, 13)):
            errs.append(f"JANコードが 8 桁 / 13 桁の数字ではありません（{jan or '空'}）")
        elif jan and not jan_check_ok(jan):
            errs.append("JANコードのチェックディジットが合いません")
        elif jan and jan in seen:
            errs.append(f"JANコードがファイル内で重複しています（{seen[jan]} 行目と同じ）")
        elif jan:
            seen[jan] = i

        code = rec[COL_ITEM_CODE] or jan
        rec[COL_ITEM_CODE] = code
        if not code or len(code) > 250 or any(ord(c) < 32 for c in code):
            errs.append("型番 / SKU は 1〜250 字で記入してください（制御文字不可）"
                        + ("。JAN が無い商品はテンプレの型番欄に手入力必須" if not jan else ""))
        elif code in seen_codes:
            errs.append(f"型番 / SKU がファイル内で重複しています（{seen_codes[code]} 行目と同じ）")
        else:
            seen_codes[code] = i
        if not jan and not rec[COL_NAME_MANUAL]:
            errs.append("JAN が無い商品はテンプレのアイテム名（任意）に商品名を手入力してください（必須）")

        cost = _num(rec[COL_COST])
        if cost is None or not cost > 0:
            errs.append(f"商品原価が正の数ではありません（{rec[COL_COST] or '空'}）")
        if rec[COL_SUPPLIER] not in suppliers:
            errs.append(f"仕入先が選択肢にありません（{rec[COL_SUPPLIER] or '空'}）")
        if rec[COL_OWNER] not in owners:
            errs.append(f"商品担当者が選択肢にありません（{rec[COL_OWNER] or '空'}）")
        if rec[COL_LARGE] not in mbl:
            errs.append(f"大分類が選択肢にありません（{rec[COL_LARGE] or '空'}）")
        elif rec[COL_MIDDLE] not in mbl[rec[COL_LARGE]]:
            errs.append(f"大分類「{rec[COL_LARGE]}」に中分類「{rec[COL_MIDDLE] or '空'}」はありません")
        if rec[COL_CARTON] == "":
            warns.append(f"{COL_CARTON} が空です（必須・手入力）")
        elif not ((n := _num(rec[COL_CARTON])) is not None and n > 0):
            errs.append(f"{COL_CARTON} が正の数ではありません（{rec[COL_CARTON]}）")
        if not rec[COL_LOT]:
            warns.append(f"{COL_LOT} が空です")
        errs.extend(validate_fields({c: v for c, v in rec.items() if c not in (*COL_PKG_DIMS, COL_PKG_WEIGHT)}, choices=choices))

        for c in (*COL_PKG_DIMS, COL_PKG_WEIGHT):
            if rec[c] and not ((n := _num(rec[c])) is not None and n > 0):
                warns.append(f"{c} が正の数ではないので空にしました（{rec[c]}）")
                rec[c] = ""
        if any(rec[c] == "" for c in COL_PKG_DIMS):
            warns.append("パッケージ 3 辺のどれかが空です（網調べ後も不明なら 1 cm の仮値・非実測）")

        if errs:
            issues.extend(Issue(i, code or jan, "error", m) for m in errs)
            continue
        issues.extend(Issue(i, code or jan, "warn", m) for m in warns)
        rows.append(rec)
    return rows, issues


def validate_fields(row: dict, *, choices=None) -> list[str]:
    """任意入力でも原本の数値・日付・文字数・選択肢条件を守る。必須の空欄は別の待入力表へ。"""
    errors = []
    values_by_key = {_norm(k): v for k, v in FIELD_CHOICES.items()}
    for col, raw in row.items():
        value = _s(raw)
        if not value:
            continue
        if col in NUMERIC_COLUMNS:
            n = _num(_nfkc(value))
            positive = "(cm)" in col or "(g)" in col or col in (COL_CARTON, COL_COST)
            if n is None or (n <= 0 if positive else n < 0):
                errors.append(f"{col} が{'正の' if positive else '0 以上の'}有限数ではありません（{value}）")
        key = _norm(col)
        if key in values_by_key and value not in values_by_key[key]:
            errors.append(f"{col} が選択肢にありません（{value}）")
        hint = NST_FIELD_HELP.get(col, "")
        if "最大60文字" in hint and len(value) > 60:
            errors.append(f"{col} は最大 60 字です")
        if "最大999文字" in hint and len(value) > 999:
            errors.append(f"{col} は最大 999 字です")
        if "全角カッコ、カンマ禁止" in hint and any(c in value for c in "（）,，"):
            errors.append(f"{col} に禁止文字（全角カッコ・カンマ）があります")
        if "ローマ数字禁止" in hint and re.search(r"[Ⅰ-ⅿ]", value):
            errors.append(f"{col} に禁止文字（ローマ数字）があります")
        if col == "生原価" and not re.fullmatch(r"(?:USD|RMB)\s*\d+(?:\.\d+)?", _nfkc(value)):
            errors.append("生原価は USD または RMB を数字の前につけてください（例 RMB53）")
        if col == "中国手数料" and not re.fullmatch(r"\d+(?:\.\d+)?%", _nfkc(value)):
            errors.append("中国手数料は数字の後に % をつけてください（例 5.3%）")
        if col == "発売日":
            try:
                if datetime.datetime.strptime(value, "%Y/%m/%d").strftime("%Y/%m/%d") != value:
                    raise ValueError(value)
            except ValueError:
                errors.append("発売日は YYYY/MM/DD で記入してください")
    if choices is not None:
        if row.get(COL_LARGE) not in choices.middle_by_large:
            errors.append("大分類が選択肢にありません")
        elif row.get(COL_MIDDLE) not in choices.middle_by_large[row[COL_LARGE]]:
            errors.append(f"大分類「{row[COL_LARGE]}」に中分類「{row.get(COL_MIDDLE, '')}」はありません")
        if row.get(COL_SUPPLIER) not in choices.suppliers:
            errors.append("仕入先が選択肢にありません")
        if row.get(COL_OWNER) not in choices.owners:
            errors.append("商品担当者が選択肢にありません")
    return errors
